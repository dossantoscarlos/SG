"""Geração dos arquivos de exportação do relatório."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lib import ui


def _rename(df: pd.DataFrame, columns: dict[str, str]) -> pd.DataFrame:
    return df.rename(columns=columns).copy()


def _currency_columns(df: pd.DataFrame) -> set[str]:
    names = {
        "Receita",
        "Total",
        "Desconto",
        "Frete",
        "Valor unitário",
        "Total do item",
        "Lucro bruto",
        "Ticket médio",
    }
    return {column for column in df.columns if column in names}


def build_excel(report: dict) -> bytes:
    summary = report["summary"]
    start = report["start"].strftime("%d/%m/%Y")
    end = report["end"].strftime("%d/%m/%Y")

    summary_df = pd.DataFrame(
        {
            "Indicador": [
                "Período",
                "Receita confirmada",
                "Pedidos confirmados",
                "Ticket médio",
                "Itens vendidos",
                "Lucro bruto estimado",
                "Margem bruta",
                "Descontos concedidos",
            ],
            "Valor": [
                f"{start} a {end}",
                ui.money(summary["revenue"]),
                ui.number_br(summary["orders_count"]),
                ui.money(summary["avg_ticket"]),
                ui.number_br(summary["units_sold"]),
                ui.money(summary["gross_profit"]),
                ui.percent(summary["gross_margin"]),
                ui.money(summary["discounts"]),
            ],
        }
    )

    daily = _rename(
        report["daily"],
        {
            "data": "Data",
            "pedidos": "Pedidos",
            "pedidos_pagos": "Pedidos confirmados",
            "receita": "Receita",
            "novos_clientes": "Novos clientes",
            "unidades": "Unidades",
        },
    )
    if "Data" in daily:
        daily["Data"] = pd.to_datetime(daily["Data"]).dt.strftime("%d/%m/%Y")

    categories = _rename(
        report["by_category"],
        {"categoria": "Categoria", "unidades": "Unidades", "receita": "Receita"},
    )
    products = _rename(
        report["top_products"],
        {
            "produto": "Produto",
            "categoria": "Categoria",
            "unidades": "Unidades",
            "receita": "Receita",
            "lucro_bruto": "Lucro bruto",
        },
    )
    customers = _rename(
        report["top_customers"],
        {
            "cliente": "Cliente",
            "email": "E-mail",
            "cidade": "Cidade",
            "uf": "UF",
            "pedidos": "Pedidos",
            "receita": "Receita",
            "ultimo_pedido": "Último pedido",
        },
    )
    if "Último pedido" in customers:
        customers["Último pedido"] = pd.to_datetime(
            customers["Último pedido"], errors="coerce"
        ).dt.strftime("%d/%m/%Y")
    payments = _rename(
        report["payment_methods"],
        {
            "forma_pagamento": "Forma de pagamento",
            "pedidos": "Pedidos",
            "receita": "Receita",
        },
    )
    orders = _rename(
        report["orders_detail"],
        {
            "pedido": "Pedido",
            "data": "Data",
            "cliente": "Cliente",
            "status": "Status",
            "pagamento": "Pagamento",
            "canal": "Canal",
            "itens": "Itens",
            "desconto": "Desconto",
            "frete": "Frete",
            "total": "Total",
        },
    )
    if "Data" in orders:
        orders["Data"] = pd.to_datetime(orders["Data"]).dt.strftime("%d/%m/%Y")
    items = _rename(
        report["items_detail"],
        {
            "pedido": "Pedido",
            "data": "Data",
            "cliente": "Cliente",
            "sku": "SKU",
            "produto": "Produto",
            "categoria": "Categoria",
            "quantidade": "Quantidade",
            "valor_unitario": "Valor unitário",
            "total_item": "Total do item",
        },
    )
    if "Data" in items:
        items["Data"] = pd.to_datetime(items["Data"]).dt.strftime("%d/%m/%Y")

    sheets = {
        "Resumo": summary_df,
        "Vendas por dia": daily,
        "Categorias": categories,
        "Top produtos": products,
        "Top clientes": customers,
        "Pagamentos": payments,
        "Pedidos": orders,
        "Itens": items,
    }

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="0F766E")
        header_font = Font(color="FFFFFF", bold=True, name="Calibri")
        body_font = Font(color="1F2937", name="Calibri")
        for sheet_name, frame in sheets.items():
            worksheet = workbook[sheet_name]
            worksheet.freeze_panes = "A2"
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")
            worksheet.row_dimensions[1].height = 24

            if sheet_name != "Resumo":
                for row in worksheet.iter_rows(min_row=2):
                    for cell in row:
                        cell.font = body_font

            for index, column in enumerate(frame.columns, start=1):
                values = [str(column)] + [str(value) for value in frame[column].head(80)]
                width = min(max(len(value) for value in values) + 2, 44)
                worksheet.column_dimensions[get_column_letter(index)].width = max(10, width)
                if column in _currency_columns(frame):
                    for row in range(2, len(frame) + 2):
                        worksheet.cell(row=row, column=index).number_format = "R$ #,##0.00"

            if len(frame.columns) > 0 and len(frame) > 0 and sheet_name != "Resumo":
                worksheet.auto_filter.ref = worksheet.dimensions

    return buffer.getvalue()


def build_csv(report: dict) -> bytes:
    orders = _rename(
        report["orders_detail"],
        {
            "pedido": "Pedido",
            "data": "Data",
            "cliente": "Cliente",
            "status": "Status",
            "pagamento": "Pagamento",
            "canal": "Canal",
            "itens": "Itens",
            "desconto": "Desconto",
            "frete": "Frete",
            "total": "Total",
        },
    )
    if "Data" in orders:
        orders["Data"] = pd.to_datetime(orders["Data"]).dt.strftime("%d/%m/%Y")
    content = orders.to_csv(index=False, sep=";", decimal=",")
    return ("\ufeff" + content).encode("utf-8")


def export_filename(report: dict, extension: str) -> str:
    start = report["start"].strftime("%Y-%m-%d")
    end = report["end"].strftime("%Y-%m-%d")
    return f"relatorio_sg_{start}_a_{end}.{extension}"
